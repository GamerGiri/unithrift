import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_callout_border(cell, color_hex="2563EB", sz="24"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def set_table_borders(table, color_hex="CBD5E1"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="{color_hex}"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_row(row, bg_hex, text_color, is_header=False, font_size=8.5):
    for cell in row.cells:
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            for run in p.runs:
                run.font.name = 'Segoe UI'
                run.font.size = Pt(font_size)
                run.font.color.rgb = text_color
                if is_header:
                    run.font.bold = True

def add_header_footer(doc):
    for section in doc.sections:
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("UniThrift: Campus Academic ReUse Marketplace | CSE 471 Documentation")
        hrun.font.name = 'Segoe UI'
        hrun.font.size = Pt(8)
        hrun.font.color.rgb = RGBColor(100, 116, 139)
        hrun.font.italic = True

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun1 = fp.add_run("Student ID: 2023200000732  |  Southeast University  |  Department of CSE")
        frun1.font.name = 'Segoe UI'
        frun1.font.size = Pt(8)
        frun1.font.color.rgb = RGBColor(100, 116, 139)

def create_callout(doc, title, text, bg_hex="EFF6FF", border_hex="2563EB", title_color=RGBColor(30, 58, 138)):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_callout_border(cell, border_hex, sz="24")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=140)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_t = p.add_run(f"{title}: ")
    run_t.font.name = 'Segoe UI'
    run_t.font.bold = True
    run_t.font.size = Pt(9)
    run_t.font.color.rgb = title_color
    
    run_b = p.add_run(text)
    run_b.font.name = 'Segoe UI'
    run_b.font.size = Pt(8.5)
    run_b.font.color.rgb = RGBColor(51, 65, 85)

print("Docx helper module verified.")
